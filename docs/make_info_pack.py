"""Build docs/Earn-to-Fire-info-pack.docx for the IBB organisers (python-docx).
Run: python docs/make_info_pack.py
"""
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "screenshots")
OUT = os.path.join(HERE, "Earn-to-Fire-info-pack.docx")
LOGO = os.path.join(HERE, "..", "site", "public", "assets", "ibb-logo.png")

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)
doc.styles["Normal"].font.name = "Segoe UI"
doc.styles["Normal"].font.size = Pt(10.5)


def para(text, bold_lead=None, size=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    return p


def bullets(items, style="List Bullet"):
    for lead, rest in items:
        p = doc.add_paragraph(style=style)
        if lead:
            p.add_run(lead).bold = True
        p.add_run(rest)


def shot(name, caption):
    p = para(caption)
    p.paragraph_format.keep_with_next = True   # never strand a caption above a page break
    p.paragraph_format.space_before = Pt(10)
    doc.add_picture(os.path.join(SHOTS, name), width=Cm(11.5))


# --- title ---------------------------------------------------------------
if os.path.exists(LOGO):
    doc.add_picture(LOGO, width=Cm(4.5))
doc.add_heading("Earn to Fire: STEMfest Newcastle info pack", level=0)
para("InfoSec Battle Bots · 3 October 2026", size=9)

doc.add_heading("What it is", level=1)
p = para("Earn to Fire is a new kids' STEM mode for the IBB arena, for ages 12 to 14. Teams answer "
         "robot-themed science, maths and coding questions on an iPad, and ")
p.add_run("every right answer fires a real arena hazard.").bold = True
p = para("It launches at ")
p.add_run("STEMfest Newcastle, 6 to 8 October 2026").bold = True
p.add_run(", for around 1,500 students in groups of about 10, roughly 50 groups a day in 7-minute slots.")

doc.add_heading("How it plays", level=1)
para("It's Red vs Blue. Each team of about five has one iPad and two bots in the arena.")
bullets([
    ("Two drivers ", "per team drive the bots, keeping their own out of the pit and pushing the other team's in."),
    ("The brains ", "(the rest of the team) tap a weapon on the iPad and answer its question."),
    ("Right answer: ", "that weapon fires for real in the arena, and the team scores its points."),
    ("Wrong answer: ", "the iPad shows the right answer and explains why. That team waits 15 seconds before trying that weapon again."),
    ("Recharge times are shared, ", "so firing a weapon also blocks the other team from it."),
    ("A round lasts 2 minutes. ", "It ends early if both of one team's bots are pitted, which is a knockout."),
], style="List Number")

doc.add_heading("Scoring", level=1)
p = para("A team's score is ")
p.add_run("STEM points plus bot points").bold = True
p.add_run(", so the brains and the drivers both win it.")
rows = [("Weapon", "Question topic", "STEM points", "What fires"),
        ("Flipper", "Forces & levers", "10", "Flipper (solenoid)"),
        ("The Pit", "Circuits & electricity", "15", "Pit opens, closes after 10 s"),
        ("Spinner 1", "Gears & motion", "20", "Spinner, 8 s, random direction"),
        ("Spinner 2", "Code & logic", "20", "Spinner, 8 s, random direction"),
        ("Spinner 3", "Sensors & robot brains", "20", "Spinner, 8 s, random direction"),
        ("MEGA SPIN", "Boss challenge", "50", "All three spinners, 10 s")]
t = doc.add_table(rows=len(rows), cols=4)
t.style = "Light Grid Accent 1"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = t.cell(i, j)
        cell.text = val
        if i == 0:
            cell.paragraphs[0].runs[0].bold = True
p = para("")
p.paragraph_format.space_before = Pt(8)
p.add_run("Bot points: ").bold = True
p.add_run("+1 every second for each of your bots still in the arena, and ")
p.add_run("+50 for every enemy bot pitted").bold = True
p.add_run(". When the clock runs out, the team with more bots still in wins the fight.")

doc.add_heading("What kids learn", level=1)
para("There are 72 KS3-level questions, 12 per weapon, all robot-themed. Examples: Ohm's law, gear "
     "ratios, binary, ultrasonic ranging, kinetic energy, and how this arena's own solenoid flipper works.")
para("Every answer, right or wrong, comes with a one-line explanation. Teachers get a take-away sheet "
     "listing the topics, with five follow-up questions for class.")

doc.add_heading("Referee tablet and safety", level=1)
p = para("One tablet runs the whole game. The arena ")
p.add_run("always starts SAFE").bold = True
p.add_run(": answers score, but nothing moves until the referee arms it.")
bullets([
    ("Round clock: ", "start, stop and reset. Default 2:00, with presets from 1 to 5 minutes."),
    ("Bot fight: ", "a PITTED button for each of the four bots, with undo. The kids can name their bots."),
    ("Standings: ", "STEM points, bot time, pit bonus and total for each team, updated live."),
    ("ALL STOP: ", "disarms, stops the clock and switches every spinner off, at any time."),
    ("Auto-stops: ", "spinners switch off after 8 seconds (MEGA SPIN 10), and the pit closes itself after 10."),
])
p = para("Rule for every volunteer: ")
p.add_run("disarm before any hand goes near the arena.").bold = True

doc.add_heading("Screenshots", level=1)
para("These were taken from a demo game, so the team and bot names are made up.")
shot("1-team-ipad.jpg", "The team iPad: both scores, the round clock, and the six weapons with their points.")
shot("2-question.jpg", "A question: each weapon asks one on its own topic.")
shot("3-correct-fires.jpg", "A right answer fires the weapon.")
shot("4-wrong-explains.jpg", "A wrong answer explains why.")
shot("5-referee-clock.jpg", "The referee tablet: ALL STOP, arming and the round clock.")
shot("6-referee-botfight.jpg", "Bot fight and standings: a PITTED button for each bot, and live totals.")

doc.add_heading("Status and links", level=1)
bullets([
    ("Done: ", "hardware tested on 2 October: every weapon, the auto-stops and ALL STOP."),
    ("Next: ", "test the team pages in Safari on the kids' iPads."),
    ("Next: ", "print and laminate the pack: poster, briefing script, weapon cards, referee sheet, leaderboard, teacher handout."),
])
para("More about the schools programme: www.infosecbattlebots.com/schools")
para("The game is open source: github.com/Ben-PyroGuy-Doch/IBB-Arena-STEM-Challenge")
para("Helping at STEMfest? Ask Ben for a walkthrough before Tuesday.")

doc.save(OUT)
print("wrote", OUT)
